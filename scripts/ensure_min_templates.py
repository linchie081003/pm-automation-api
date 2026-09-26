"""Create minimal placeholder templates when org files are not yet copied."""

from pathlib import Path

from openpyxl import Workbook
from pptx import Presentation

from app.config import settings
from app.services.templates.loader import load_manifest


def main() -> None:
    root = settings.templates_path
    root.mkdir(parents=True, exist_ok=True)
    manifest_path = root / "manifest.json"
    if not manifest_path.is_file():
        manifest_path.write_text(
            '{"weekly_report.xlsx":"weekly_report.xlsx","weekly_report.pptx":"weekly_report.pptx",'
            '"pre_kickoff_deck.pptx":"pre_kickoff_deck.pptx","kickoff_deck.pptx":"kickoff_deck.pptx",'
            '"task_export.xlsx":"task_export.xlsx"}',
            encoding="utf-8",
        )
    manifest = load_manifest()
    for _key, filename in manifest.items():
        path = root / filename
        if path.is_file():
            continue
        if filename.endswith(".xlsx"):
            wb = Workbook()
            ws = wb.active
            ws.title = "Data"
            ws["A1"] = "{{PROJECT_NAME}}"
            ws["A2"] = "{{PROJECT_CODE}}"
            wb.save(path)
        elif filename.endswith(".pptx"):
            prs = Presentation()
            blank = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
            slide = prs.slides.add_slide(blank)
            box = slide.shapes.add_textbox(0, 0, prs.slide_width, prs.slide_height)
            box.text_frame.text = "{{PROJECT_NAME}}\n{{SCOPE}}\n{{MILESTONE_TABLE}}"
            prs.save(path)
    print(f"Ensured templates under {root}")


if __name__ == "__main__":
    main()
