"""Inspect yyyymmdd-Template.xlsx sheets (run from backend dir)."""
from pathlib import Path

from openpyxl import load_workbook

TEMPLATE = Path(__file__).resolve().parents[1] / "templates" / "yyyymmdd-Template.xlsx"


def main():
    if not TEMPLATE.is_file():
        print(f"Missing {TEMPLATE}")
        return
    wb = load_workbook(TEMPLATE, read_only=True, data_only=True)
    print("Sheets:", wb.sheetnames)
    for name in wb.sheetnames[:6]:
        ws = wb[name]
        print(f"\n=== {name} ===")
        for i, row in enumerate(ws.iter_rows(max_row=8, max_col=6, values_only=True)):
            print(i + 1, row)
    wb.close()


if __name__ == "__main__":
    main()
