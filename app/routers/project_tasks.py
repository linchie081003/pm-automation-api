import uuid
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from openpyxl import Workbook
from sqlalchemy.orm import Session

from app.config import settings
from app.core.deps import get_current_user, get_permission_codes
from app.core.project_access import ensure_permission, ensure_project_read
from app.database import get_db
from app.models import Project, User
from app.services.activity import log_activity
from app.services.clickup import ClickUpSyncError, sync_project_tasks, task_recap
from app.services.templates.loader import copy_template
from app.services.templates.placeholders import build_mapping, replace_in_xlsx

router = APIRouter(prefix="/projects/{project_id}/tasks", tags=["tasks"])


@router.get("/recap")
def get_recap(
    project_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    codes: set[str] = Depends(get_permission_codes),
):
    ensure_permission(codes, "tasks.read", "tasks.export", "integrations.clickup.project")
    ensure_project_read(project_id, user, codes, db)
    return task_recap(db, project_id)


@router.post("/sync")
def sync_tasks(
    project_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    codes: set[str] = Depends(get_permission_codes),
):
    ensure_permission(codes, "integrations.clickup.project")
    ensure_project_read(project_id, user, codes, db)
    project = db.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Not found")
    try:
        result = sync_project_tasks(db, project)
        log_activity(
            db,
            project_id,
            user.id,
            "clickup.synced",
            {
                "tasks": result.get("count") or result.get("synced") or result.get("updated"),
                "detail_keys": list(result.keys())[:8],
            },
        )
        db.commit()
        return result
    except ClickUpSyncError as e:
        raise HTTPException(status_code=400, detail=str(e)) from e
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Sync ClickUp gagal: {e}") from e


@router.get("/export")
def export_tasks(
    project_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    codes: set[str] = Depends(get_permission_codes),
):
    ensure_permission(codes, "tasks.export")
    ensure_project_read(project_id, user, codes, db)
    project = db.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Not found")
    recap = task_recap(db, project_id)
    out_dir = Path(settings.upload_dir) / "exports" / str(project_id)
    out_dir.mkdir(parents=True, exist_ok=True)
    dest = out_dir / f"tasks_{project.code}_{uuid.uuid4().hex[:8]}.xlsx"
    mapping = build_mapping(
        PROJECT_NAME=project.name,
        PROJECT_CODE=project.code,
        TASK_TABLE="\n".join(f"- {t['name']}" for t in recap["tasks"]),
    )
    try:
        copy_template("task_export.xlsx", dest)
        replace_in_xlsx(dest, mapping)
        from openpyxl import load_workbook

        wb = load_workbook(str(dest))
        ws = wb.active
        ws.append(["Name", "Status", "Due", "Closed", "URL"])
        for t in recap["tasks"]:
            ws.append([t["name"], t["status"], t["due_date"], t["is_closed"], t["url"]])
        wb.save(str(dest))
    except FileNotFoundError:
        wb = Workbook()
        ws = wb.active
        ws.title = "Tasks"
        ws.append(["Name", "Status", "Due", "Closed", "URL"])
        for t in recap["tasks"]:
            ws.append([t["name"], t["status"], t["due_date"], t["is_closed"], t["url"]])
        wb.save(dest)
    return FileResponse(dest, filename=dest.name)
