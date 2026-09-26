from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation


def replace_in_pptx(path: Path, mapping: dict[str, str]) -> None:
    prs = Presentation(str(path))
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    text = run.text
                    for key, val in mapping.items():
                        text = text.replace(key, val or "")
                    run.text = text
    prs.save(str(path))


def replace_in_xlsx(path: Path, mapping: dict[str, str]) -> None:
    wb = load_workbook(str(path))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    val = cell.value
                    for key, repl in mapping.items():
                        val = val.replace(key, repl or "")
                    cell.value = val
    wb.save(str(path))


def build_mapping(**kwargs: str) -> dict[str, str]:
    return {f"{{{{{k}}}}}": (v or "") for k, v in kwargs.items()}
