"""Fill yyyymmdd-Template.xlsx (S-curve + milestone bar chart sheets) from PDC data."""
from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    ClickUpTaskCache,
    IntegrationSettings,
    Milestone,
    ProgressSnapshot,
    ProgressSnapshotSource,
    Project,
    ProjectHealthConfig,
    ProjectPo,
    ProjectSph,
    TimelineItemType,
    WeeklyReport,
)
from app.services.progress import build_clickup_lookups, clickup_status_mapping_context
from app.services.timeline_report_metrics import (
    ReportRowProgressContext,
    display_row_actual_pct,
    display_row_duration_days,
    display_row_planned_fraction,
    display_row_planned_week_fraction,
    display_row_weight_fraction,
    include_in_report_timeline_list,
)
from app.services.report_calendar import (
    extend_anchors_for_project_end,
    period_for_report_date,
    weekly_report_dates,
)
from app.services.schedule import compute_spi, snapshot_key_for_report_date
from app.services.schedule_window import (
    kickoff_milestones,
    project_report_end_date,
    project_report_start_date,
    schedule_bounds,
)
from app.services.templates.loader import copy_template
from app.services.timeline_display import timeline_display_rows_for_project
from app.core.timezone import today_jakarta

_SCURVE_ROMAN = (
    "I",
    "II",
    "III",
    "IV",
    "V",
    "VI",
    "VII",
    "VIII",
    "IX",
    "X",
    "XI",
    "XII",
)

_SHEETS_TO_DROP = ("Catatan Perbaikan", "Info Mingguan", "SCurve D1")

_PHASE_FONT = Font(bold=True, size=13)
_TASK_FONT = Font(size=11)

_SCURVE_FIRST_WEEK_COL = 16  # P
_SCURVE_LAST_WEEK_COL = 63  # BK
_SCURVE_TYPE_COL = 14  # N — label tipe (phase / milestone / task / subtask)
_LOG_MINGGUAN_FIRST_WEEK_ROW = 8
_LOG_MINGGUAN_LAST_WEEK_ROW = 55
_SCURVE_CLEAR_B74_E112 = (74, 112, 2, 5)  # rows, cols B–E


def _as_date(d: date | datetime | None) -> datetime | None:
    if d is None:
        return None
    if isinstance(d, datetime):
        return d.replace(hour=0, minute=0, second=0, microsecond=0)
    return datetime(d.year, d.month, d.day)


def _parse_row_date(val: object) -> date | None:
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    try:
        return date.fromisoformat(str(val)[:10])
    except ValueError:
        return None


def _row_actual_fraction(row: dict, ctx: ReportRowProgressContext, as_of: date) -> float:
    return display_row_actual_pct(ctx, row, as_of) / 100.0


def _row_dates(row: dict) -> tuple[date | None, date | None]:
    start = _parse_row_date(row.get("display_start") or row.get("start_date"))
    end = _parse_row_date(row.get("display_end") or row.get("target_date"))
    return start, end


def _row_display_name(row: dict) -> str:
    return (row.get("name") or row.get("clickup_name") or "").strip()


def _excel_timeline_rows(display: list[dict]) -> list[tuple[str | None, dict, bool]]:
    """(Roman numeral, row, is_phase) in timeline order; skips unnamed rows."""
    out: list[tuple[str | None, dict, bool]] = []
    phase_i = 0
    for row in display:
        if not _row_display_name(row):
            continue
        it = row.get("item_type") or ""
        if it == TimelineItemType.phase.value:
            roman = _SCURVE_ROMAN[phase_i] if phase_i < len(_SCURVE_ROMAN) else str(phase_i + 1)
            phase_i += 1
            out.append((roman, row, True))
        else:
            out.append((None, row, False))
    return out


def _scurve_type_label(row: dict) -> str:
    it = row.get("item_type") or ""
    if it == TimelineItemType.phase.value:
        return "Phase"
    if it == TimelineItemType.milestone.value:
        return "Milestone"
    if it == TimelineItemType.subtask.value:
        return "Subtask"
    if it == TimelineItemType.task.value:
        return "Task"
    return str(it).capitalize() if it else ""


def _write_scurve_row_metrics_static(
    ws,
    r: int,
    *,
    row: dict,
    weight_frac: float,
    cut_off: date,
    week_periods: list[tuple[date, date]],
    ctx: ReportRowProgressContext,
) -> None:
    """Planned/actual from Timeline Engine — static values, no NETWORKDAYS formulas."""
    planned_frac = display_row_planned_fraction(row, cut_off, ctx.db)
    actual_frac = _row_actual_fraction(row, ctx, cut_off)
    dur = display_row_duration_days(row, ctx.db)
    ws.cell(r, 8).value = dur
    ws.cell(r, 9).value = round(planned_frac, 6)
    ws.cell(r, 10).value = round(weight_frac * planned_frac, 6)
    ws.cell(r, 11).value = round(actual_frac, 6)
    ws.cell(r, 12).value = round(weight_frac * (actual_frac - planned_frac), 6)
    ws.cell(r, 13).value = round(weight_frac * actual_frac, 6)
    ws.cell(r, 15).value = round(weight_frac * planned_frac, 6)
    for wi, (period_start, period_end) in enumerate(week_periods):
        col = _SCURVE_FIRST_WEEK_COL + wi
        ws.cell(r, col).value = round(
            display_row_planned_week_fraction(
                row, weight_frac, period_start, period_end, ctx.db
            ),
            6,
        )


def _pct_as_excel_fraction(value: float) -> float:
    if value > 1.0:
        return round(value / 100.0, 6)
    return round(float(value), 6)


def _snapshot_has_saved_actual(snap: ProgressSnapshot) -> bool:
    if snap.weekly_report_id is not None:
        return True
    return snap.source in (
        ProgressSnapshotSource.weekly_report,
        ProgressSnapshotSource.manual_save,
    )


def _weekly_report_actual_map(
    db: Session,
    project: Project,
    project_start: date | None,
) -> dict[date, float]:
    """
    Actual kumulatif per minggu laporan — hanya dari WeeklyReport yang sudah di-generate
    (frozen_metrics), keyed by week_start / normalized report_date.
    """
    reports = db.scalars(
        select(WeeklyReport)
        .where(WeeklyReport.project_id == project.id)
        .order_by(WeeklyReport.week_start)
    ).all()
    out: dict[date, float] = {}
    for wr in reports:
        fm = wr.frozen_metrics or {}
        raw = fm.get("actual_pct")
        if raw is None:
            continue
        frac = _pct_as_excel_fraction(float(raw))
        out[wr.week_start] = frac
        norm, _ = snapshot_key_for_report_date(
            project, wr.week_start, project_start_date=project_start
        )
        if norm not in out:
            out[norm] = frac
    return out


def _snapshot_actual_map_exact(
    db: Session,
    project: Project,
    project_start: date | None,
) -> dict[date, float]:
    """Fallback: snapshot dengan weekly report / manual save, exact week_start only."""
    snaps = db.scalars(
        select(ProgressSnapshot).where(ProgressSnapshot.project_id == project.id)
    ).all()
    out: dict[date, float] = {}
    for s in snaps:
        if not _snapshot_has_saved_actual(s):
            continue
        frac = _pct_as_excel_fraction(s.actual_cumulative_pct)
        out[s.week_start] = frac
        norm, _ = snapshot_key_for_report_date(
            project, s.week_start, project_start_date=project_start
        )
        if norm not in out:
            out[norm] = frac
    return out


def _actual_kumulatif_for_week(
    report_actuals: dict[date, float],
    snap_actuals: dict[date, float],
    project: Project,
    anchor: date,
    project_start: date | None,
) -> float | None:
    norm, _ = snapshot_key_for_report_date(
        project, anchor, project_start_date=project_start
    )
    for key in (norm, anchor):
        if key in report_actuals:
            return report_actuals[key]
    for key in (norm, anchor):
        if key in snap_actuals:
            return snap_actuals[key]
    return None


def _week_index_for_status(
    status_date: date,
    anchors: list[date],
    project_start: date | None,
) -> int:
    if not anchors:
        return 1
    for i, end in enumerate(anchors):
        if i == 0:
            start = project_start or end
        else:
            start = anchors[i - 1] + timedelta(days=1)
        if start <= status_date <= end:
            return i + 1
    if project_start and status_date < project_start:
        return 1
    return len(anchors)


def _clear_scurve_cells_b74_e112(ws) -> None:
    r0, r1, c0, c1 = _SCURVE_CLEAR_B74_E112
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            ws.cell(r, c).value = None


def _remove_sheet_table_styles(ws) -> None:
    if ws is None:
        return
    for name in list(getattr(ws, "tables", {}) or {}):
        del ws.tables[name]
    ws.auto_filter.ref = None
    max_r = max(ws.max_row or 1, 500)
    for row in ws.iter_rows(min_row=1, max_row=max_r, min_col=1, max_col=20):
        for cell in row:
            cell.fill = PatternFill(fill_type=None)
            cell.border = Border()
            cell.style = "Normal"


def _trim_workbook_sheets(wb) -> None:
    for name in _SHEETS_TO_DROP:
        if name in wb.sheetnames:
            del wb[name]
    if "Sheet2" in wb.sheetnames:
        wb["Sheet2"].title = "Task"
    elif "Task" not in wb.sheetnames:
        wb.create_sheet("Task")


def _weekly_report_anchor_list(
    db: Session,
    project: Project,
) -> tuple[list[date], date | None]:
    pstart = project_report_start_date(db, project)
    pend, _ = project_report_end_date(db, project.id)
    if not pend:
        sch_start, sch_end = schedule_bounds(db, project.id)
        pend = sch_end or sch_start
    if not pend:
        return [], pstart
    anchors, _ = weekly_report_dates(
        pstart,
        pend,
        project.weekly_report_anchor_weekday,
        project.weekly_report_first_anchor_date,
        cap_at=pend,
    )
    anchors = extend_anchors_for_project_end(
        anchors,
        pend,
        project.weekly_report_cutoff_offset_days,
    )
    return anchors, pstart


def export_yyyymmdd_workbook(
    dest: Path,
    *,
    db: Session,
    project: Project,
    milestones: list[Milestone],
    anchor: date,
    cut_off: date,
    planned_pct: float,
    actual_pct: float,
    spi: float | None = None,
) -> Path:
    copy_template("yyyymmdd-Template.xlsx", dest)
    wb = load_workbook(dest)
    _trim_workbook_sheets(wb)

    display = timeline_display_rows_for_project(db, project.id)
    ms_list = kickoff_milestones(db, project.id) or milestones
    ms_by_id = {m.id: m for m in ms_list}
    caches = list(
        db.scalars(
            select(ClickUpTaskCache).where(ClickUpTaskCache.project_id == project.id)
        ).all()
    )
    excel_rows = _excel_timeline_rows(display)
    weight_fractions = [display_row_weight_fraction(row) for _, row, _ in excel_rows]

    sph = db.get(ProjectSph, project.id)
    po = db.get(ProjectPo, project.id)

    anchors, project_start = _weekly_report_anchor_list(db, project)
    if not anchors:
        anchors = [anchor]
    if not project_start:
        project_start = project_report_start_date(db, project)

    spi_val = spi
    if spi_val is None and planned_pct > 0:
        spi_val = compute_spi(actual_pct, planned_pct)
    deviation_pp = actual_pct - planned_pct

    scurve_data_start = 21
    with clickup_status_mapping_context(db):
        tasks_by_id, milestone_cache = build_clickup_lookups(ms_list, caches)
        progress_ctx = ReportRowProgressContext(
            db=db,
            project=project,
            milestones=ms_list,
            ms_by_id=ms_by_id,
            tasks_by_id=tasks_by_id,
            milestone_cache=milestone_cache,
            caches=caches,
        )
        phase_scurve_rows = _fill_scurve_sheet(
            wb["SCurve"],
            project=project,
            sph=sph,
            po=po,
            excel_rows=excel_rows,
            weight_fractions=weight_fractions,
            progress_ctx=progress_ctx,
            anchor=anchor,
            cut_off=cut_off,
            anchors=anchors,
            project_start=project_start,
            data_start=scurve_data_start,
        )
    health_cfg = db.get(ProjectHealthConfig, project.id)
    _fill_log_mingguan(
        wb["Log_Mingguan"] if "Log_Mingguan" in wb.sheetnames else None,
        db=db,
        project=project,
        cut_off=cut_off,
        planned_pct=planned_pct,
        actual_pct=actual_pct,
        deviation_pp=deviation_pp,
        spi=spi_val,
        spi_green_min=health_cfg.spi_green_min if health_cfg else 1.0,
        spi_yellow_min=health_cfg.spi_yellow_min if health_cfg else 0.9,
        anchors=anchors,
        project_start=project_start,
        week_col_count=min(
            len(anchors),
            _SCURVE_LAST_WEEK_COL - _SCURVE_FIRST_WEEK_COL + 1,
        ),
    )
    _fill_milestone_progress_sheet(
        wb["milestone progress"] if "milestone progress" in wb.sheetnames else None,
        phase_scurve_rows=phase_scurve_rows,
        progress_ctx=progress_ctx,
        cut_off=cut_off,
    )
    _fill_libur_sheet(wb["Libur"] if "Libur" in wb.sheetnames else None, db)
    _fill_task_sheet(
        wb["Task"] if "Task" in wb.sheetnames else None,
        excel_rows,
        weight_fractions,
        progress_ctx=progress_ctx,
        cut_off=cut_off,
    )

    wb.save(dest)
    wb.close()
    return dest


def _fill_scurve_sheet(
    ws,
    *,
    project: Project,
    sph: ProjectSph | None,
    po: ProjectPo | None,
    excel_rows: list[tuple[str | None, dict, bool]],
    weight_fractions: list[float],
    progress_ctx: ReportRowProgressContext,
    anchor: date,
    cut_off: date,
    anchors: list[date],
    project_start: date | None,
    data_start: int,
) -> list[tuple[str, dict, int]]:
    """Returns phase rows as (roman, row_dict, scurve_excel_row)."""
    ws.cell(18, _SCURVE_TYPE_COL).value = "Tipe"

    ws["A20"].value = None
    ws["B20"].value = None
    for col in range(3, 14):
        ws.cell(20, col).value = None

    period_len = int(project.weekly_report_cutoff_offset_days or 6)
    ws["C2"] = project.name
    ws["C3"] = project.client_name or (sph.sph_client if sph else "")
    ws["C4"] = (sph.pic_user_name if sph and sph.pic_user_name else "") or ""
    ws["C5"] = _as_date(project.planned_end_date)
    ws["C6"] = _as_date(po.po_due_date if po and po.po_due_date else project.po_due_date)
    ws["C7"] = _as_date(anchor)
    ws["C8"] = _as_date(cut_off)

    week_cols = min(len(anchors), _SCURVE_LAST_WEEK_COL - _SCURVE_FIRST_WEEK_COL + 1)
    week_periods: list[tuple[date, date]] = []
    for i in range(week_cols):
        col = _SCURVE_FIRST_WEEK_COL + i
        report_date = anchors[i]
        if i == 0:
            row16 = project_start or report_date
        else:
            row16 = anchors[i - 1] + timedelta(days=1)
        _, period_end = period_for_report_date(
            report_date,
            period_len,
            project_start,
            report_weekday=project.weekly_report_anchor_weekday,
            explicit_first_report_date=project.weekly_report_first_anchor_date,
        )
        row17 = period_end or report_date
        week_periods.append((row16, row17))
        ws.cell(16, col).value = _as_date(row16)
        ws.cell(17, col).value = _as_date(row17)
        ws.cell(18, col).value = f"Week-{i + 1:02d}"

    max_row = 60
    for r in range(data_start, max_row + 1):
        for col in range(1, _SCURVE_LAST_WEEK_COL + 1):
            ws.cell(r, col).value = None

    phase_row_nums: list[int] = []
    phase_meta: list[tuple[str, dict, int]] = []
    r = data_start
    idx = 0

    def _write_detail_row(cr: int, child_idx: int) -> None:
        """Task / subtask / milestone — tampil saja, tanpa formula mingguan."""
        _, child, _ = excel_rows[child_idx]
        w_child = weight_fractions[child_idx] if child_idx < len(weight_fractions) else 0.0
        cname = _row_display_name(child)
        cstart, cend = _row_dates(child)
        depth = int(child.get("depth") or 0)
        ws.cell(cr, 2).value = f"{'    ' * max(1, depth)}{cname}"
        ws.cell(cr, 2).font = _TASK_FONT
        ws.cell(cr, 2).alignment = Alignment(indent=max(0, depth))
        ws.cell(cr, 3).value = round(w_child, 6)
        ws.cell(cr, 4).value = _as_date(cstart)
        ws.cell(cr, 5).value = _as_date(cend)
        ws.cell(cr, _SCURVE_TYPE_COL).value = _scurve_type_label(child)
        it = child.get("item_type") or ""
        if it != TimelineItemType.milestone.value:
            ws.cell(cr, 11).value = _row_actual_fraction(child, progress_ctx, cut_off)

    while idx < len(excel_rows) and r <= max_row:
        roman, row, is_phase = excel_rows[idx]
        weight_frac = weight_fractions[idx] if idx < len(weight_fractions) else 0.0
        name = _row_display_name(row)
        start, end = _row_dates(row)

        if is_phase:
            child_indices: list[int] = []
            j = idx + 1
            while j < len(excel_rows) and not excel_rows[j][2]:
                _, child_row, _ = excel_rows[j]
                if include_in_report_timeline_list(child_row):
                    child_indices.append(j)
                j += 1

            phase_r = r
            ws.cell(phase_r, 1).value = roman
            phase_cell = ws.cell(phase_r, 2)
            phase_cell.value = name
            phase_cell.font = _PHASE_FONT
            ws.cell(phase_r, 3).value = round(float(row.get("weight_pct") or 0) / 100.0, 6)
            ws.cell(phase_r, 4).value = _as_date(start)
            ws.cell(phase_r, 5).value = _as_date(end)
            ws.cell(phase_r, _SCURVE_TYPE_COL).value = "Phase"
            _write_scurve_row_metrics_static(
                ws,
                phase_r,
                row=row,
                weight_frac=float(row.get("weight_pct") or 0) / 100.0,
                cut_off=cut_off,
                week_periods=week_periods,
                ctx=progress_ctx,
            )

            first_child_r = phase_r + 1
            for k, child_idx in enumerate(child_indices):
                cr = first_child_r + k
                if cr > max_row:
                    break
                _write_detail_row(cr, child_idx)

            phase_row_nums.append(phase_r)
            phase_meta.append((roman or "", row, phase_r))
            r = first_child_r + len(child_indices)
            idx = j
            continue

        if not include_in_report_timeline_list(row):
            idx += 1
            continue

        ws.cell(r, 2).value = name
        ws.cell(r, 2).font = _TASK_FONT
        ws.cell(r, 3).value = round(weight_frac, 6)
        ws.cell(r, 4).value = _as_date(start)
        ws.cell(r, 5).value = _as_date(end)
        ws.cell(r, _SCURVE_TYPE_COL).value = _scurve_type_label(row)
        if (row.get("item_type") or "") != TimelineItemType.milestone.value:
            ws.cell(r, 11).value = _row_actual_fraction(row, progress_ctx, cut_off)
        r += 1
        idx += 1

    if phase_row_nums:
        parts = ",".join(f"C{n}" for n in phase_row_nums)
        ws.cell(62, 3).value = f"=SUM({parts})"
        jparts = ",".join(f"J{n}" for n in phase_row_nums)
        lparts = ",".join(f"L{n}" for n in phase_row_nums)
        mparts = ",".join(f"M{n}" for n in phase_row_nums)
        oparts = ",".join(f"O{n}" for n in phase_row_nums)
        ws.cell(62, 10).value = f"=SUM({jparts})"
        ws.cell(62, 12).value = f"=SUM({lparts})"
        ws.cell(62, 13).value = f"=SUM({mparts})"
        ws.cell(62, 15).value = f"=SUM({oparts})"
        for wi in range(week_cols):
            col = _SCURVE_FIRST_WEEK_COL + wi
            wparts = ",".join(f"{get_column_letter(col)}{n}" for n in phase_row_nums)
            ws.cell(62, col).value = f"=SUM({wparts})"

    _clear_scurve_cells_b74_e112(ws)
    return phase_meta


def _fill_libur_sheet(ws, db: Session) -> None:
    if ws is None:
        return
    settings = db.scalars(select(IntegrationSettings).limit(1)).first()
    holidays = sorted(
        list(settings.holiday_dates or []) if settings else [],
        key=lambda h: str(h.get("date") or ""),
    )
    max_r = max(ws.max_row or 1, 500)
    for r in range(1, max_r + 1):
        for c in range(1, 16):
            ws.cell(r, c).value = None
    ws.cell(1, 1).value = "Tanggal Libur"
    for i, h in enumerate(holidays):
        r = 2 + i
        d = _parse_row_date(h.get("date"))
        ws.cell(r, 1).value = _as_date(d) if d else None


def _fill_task_sheet(
    ws,
    excel_rows: list[tuple[str | None, dict, bool]],
    weight_fractions: list[float],
    *,
    progress_ctx: ReportRowProgressContext,
    cut_off: date,
) -> None:
    if ws is None:
        return
    headers = (
        "No",
        "Task",
        "Bobot",
        "Start Date",
        "Finish Date",
        "Tipe",
        "Phase",
        "Status",
        "Progress %",
    )
    for c, h in enumerate(headers, start=1):
        ws.cell(1, c).value = h
    for r in range(2, (ws.max_row or 2) + 200):
        for c in range(1, len(headers) + 1):
            ws.cell(r, c).value = None

    for idx, (roman, row, is_phase) in enumerate(excel_rows):
        r = 2 + idx
        name = _row_display_name(row)
        depth = int(row.get("depth") or 0)
        if is_phase:
            task_label = name
        else:
            task_label = f"{'    ' * max(1, depth)}{name}"
        start, end = _row_dates(row)
        ws.cell(r, 1).value = roman
        ws.cell(r, 2).value = task_label
        w_frac = weight_fractions[idx] if idx < len(weight_fractions) else 0.0
        ws.cell(r, 3).value = round(display_row_weight_fraction(row), 6)
        ws.cell(r, 4).value = _as_date(start)
        ws.cell(r, 5).value = _as_date(end)
        ws.cell(r, 6).value = _scurve_type_label(row)
        ws.cell(r, 7).value = row.get("phase_name") or ""
        ws.cell(r, 8).value = row.get("clickup_status") or row.get("status") or ""
        ws.cell(r, 9).value = round(_row_actual_fraction(row, progress_ctx, cut_off), 4)
        if is_phase:
            ws.cell(r, 2).font = _PHASE_FONT
        else:
            ws.cell(r, 2).font = _TASK_FONT

    _remove_sheet_table_styles(ws)


def _fill_log_mingguan(
    ws,
    *,
    db: Session,
    project: Project,
    cut_off: date,
    planned_pct: float,
    actual_pct: float,
    deviation_pp: float,
    spi: float | None,
    spi_green_min: float,
    spi_yellow_min: float,
    anchors: list[date],
    project_start: date | None,
    week_col_count: int,
) -> None:
    if ws is None:
        return
    ws["A4"] = _as_date(cut_off)
    if week_col_count > 0:
        c0 = get_column_letter(_SCURVE_FIRST_WEEK_COL)
        c1 = get_column_letter(_SCURVE_FIRST_WEEK_COL + week_col_count - 1)
        ws["B4"] = (
            f"=IF(SUMPRODUCT((SCurve!$C$8>={c0}$16:{c1}$16)"
            f"*(SCurve!$C$8<={c0}$17:{c1}$17))=0,"
            f"{_week_index_for_status(cut_off, anchors, project_start)},"
            f"SUMPRODUCT((SCurve!$C$8>={c0}$16:{c1}$16)"
            f"*(SCurve!$C$8<={c0}$17:{c1}$17)))"
        )
    else:
        ws["B4"] = _week_index_for_status(cut_off, anchors, project_start)
    ws["C4"] = planned_pct / 100.0 if planned_pct > 1 else planned_pct
    ws["D4"] = actual_pct / 100.0 if actual_pct > 1 else actual_pct
    ws["E4"] = deviation_pp / 100.0 if abs(deviation_pp) > 1 else deviation_pp
    ws["F4"] = spi if spi is not None else ""
    ws["K4"] = spi_green_min
    ws["L4"] = spi_yellow_min

    report_actuals = _weekly_report_actual_map(db, project, project_start)
    snap_actuals = _snapshot_actual_map_exact(db, project, project_start)
    for r in range(_LOG_MINGGUAN_FIRST_WEEK_ROW, _LOG_MINGGUAN_LAST_WEEK_ROW + 1):
        ws.cell(r, 4).value = None
        ws.cell(r, 9).value = None
    for i, anchor_date in enumerate(anchors):
        r = _LOG_MINGGUAN_FIRST_WEEK_ROW + i
        if r > _LOG_MINGGUAN_LAST_WEEK_ROW:
            break
        # Minggu setelah cut-off laporan ini: Actual Kumulatif kosong (meski ada data di DB).
        if anchor_date > cut_off:
            continue
        actual_frac = _actual_kumulatif_for_week(
            report_actuals,
            snap_actuals,
            project,
            anchor_date,
            project_start,
        )
        if actual_frac is not None:
            ws.cell(r, 4).value = actual_frac


def _fill_milestone_progress_sheet(
    ws,
    *,
    phase_scurve_rows: list[tuple[str, dict, int]],
    progress_ctx: ReportRowProgressContext,
    cut_off: date,
) -> None:
    if ws is None:
        return
    max_rows = 40
    for r in range(2, 2 + max_rows):
        for c in range(1, 14):
            ws.cell(r, c).value = None

    for idx, (roman, row, sc_row) in enumerate(phase_scurve_rows[:max_rows]):
        r = 2 + idx
        name = _row_display_name(row)
        start, end = _row_dates(row)
        ws.cell(r, 1).value = roman
        name_cell = ws.cell(r, 2)
        name_cell.value = name
        name_cell.font = _PHASE_FONT
        ws.cell(r, 3).value = round(display_row_weight_fraction(row), 6)
        ws.cell(r, 4).value = _as_date(start)
        ws.cell(r, 5).value = _as_date(end)
        dur = display_row_duration_days(row, progress_ctx.db)
        ws.cell(r, 8).value = dur
        planned_frac = display_row_planned_fraction(row, cut_off, progress_ctx.db)
        actual_frac = _row_actual_fraction(row, progress_ctx, cut_off)
        ws.cell(r, 9).value = round(planned_frac, 6)
        ws.cell(r, 10).value = round(planned_frac, 6)
        ws.cell(r, 12).value = round(actual_frac, 6)
        target_w = display_row_weight_fraction(row)
        ws.cell(r, 13).value = round(target_w * (actual_frac - planned_frac), 6)


def export_scurve_from_project(
    dest: Path,
    db: Session,
    project: Project,
    milestones: list[Milestone],
    series: list[dict],
) -> Path:
    """S-curve export endpoint — same template as weekly report snapshot."""
    last = series[-1] if series else {}
    cut_str = last.get("cut_off")
    cut_off = date.fromisoformat(cut_str) if cut_str else today_jakarta()
    anchor_str = last.get("date")
    anchor = date.fromisoformat(anchor_str) if anchor_str else cut_off
    planned = float(last.get("planned_pct") or 0)
    return export_yyyymmdd_workbook(
        dest,
        db=db,
        project=project,
        milestones=milestones,
        anchor=anchor,
        cut_off=cut_off,
        planned_pct=planned,
        actual_pct=0.0,
    )
