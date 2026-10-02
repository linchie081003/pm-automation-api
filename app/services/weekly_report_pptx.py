"""Build 4-slide weekly report deck (title, overview, progress+S-curve, timeline)."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import settings
from app.models import ClickUpTaskCache, Milestone, Project, ProjectPo, ProjectSph, WeeklyReport
from app.services.report_calendar import resolve_report_period
from app.services.schedule import scurve_points
from app.services.progress import build_clickup_lookups, clickup_status_mapping_context
from app.services.schedule_window import kickoff_milestones, project_report_start_date, schedule_bounds
from app.services.timeline_display import timeline_display_rows_for_project
from app.services.timeline_report_metrics import (
    ReportRowProgressContext,
    display_row_actual_pct,
    include_in_report_timeline_list,
)
from app.core.timezone import today_jakarta

COLOR_ORANGE = RGBColor(0xF5, 0x82, 0x20)
COLOR_NAVY = RGBColor(0x1E, 0x3A, 0x5F)
COLOR_GREEN = RGBColor(0x00, 0xA6, 0x51)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_DARK = RGBColor(0x1E, 0x29, 0x3B)
COLOR_MUTED = RGBColor(0x64, 0x74, 0x8B)
COLOR_PEACH = RGBColor(0xFF, 0xF0, 0xE6)
COLOR_PEACH_LIGHT = RGBColor(0xFF, 0xF7, 0xF0)

FONT_FACE = "Arial Nova Cond"
TITLE_SIZE = 26
KPI_LABEL_SIZE = 12
KPI_VALUE_SIZE = 22
BODY_SIZE = 9
TABLE_HEADER_FONT = 14
TABLE_BODY_FONT = 12
TABLE_HEADER_ROW_PT = 26
TABLE_DATA_ROW_PT = 22
TABLE_CELL_MARGIN_PT = 3

WEEKDAY_EN = (
    "Monday",
    "Tuesday",
    "Wednesday",
    "Thursday",
    "Friday",
    "Saturday",
    "Sunday",
)
MONTH_EN = (
    "",
    "January",
    "February",
    "March",
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
)


def _templates_root() -> Path:
    return Path(settings.templates_path)


def _title_bg_path() -> Path | None:
    for name in ("slide1_bg.png", "slide1_bg.jpg"):
        p = _templates_root() / "weekly_report" / name
        if p.is_file():
            return p
    return None


def _fmt_date_en_long(d: date | None) -> str:
    if not d:
        return "—"
    return f"{WEEKDAY_EN[d.weekday()]}, {d.day:02d} {MONTH_EN[d.month]} {d.year}"


def _fmt_date_slash(d: date | None) -> str:
    if not d:
        return "—"
    return d.strftime("%d/%m/%Y")


def _parse_iso(raw: str | None) -> date | None:
    if not raw:
        return None
    try:
        return date.fromisoformat(raw[:10])
    except ValueError:
        return None


def _point_anchor_date(pt: dict) -> date | None:
    rd = pt.get("date") or pt.get("anchor_date") or pt.get("week_start")
    return _parse_iso(rd if isinstance(rd, str) else None)


def _actual_pct_for_report_export(pt: dict, report_cut_off: date) -> float | None:
    """Actual kumulatif hanya untuk minggu s.d. cut-off laporan yang di-export."""
    d = _point_anchor_date(pt)
    if d and d > report_cut_off:
        return None
    raw = pt.get("actual_pct")
    if raw is None:
        return None
    return float(raw)


def _resolve_brief(project: Project, sph: ProjectSph | None) -> str:
    if project.project_brief and project.project_brief.strip():
        return project.project_brief.strip()
    if sph and sph.scope_text and sph.scope_text.strip():
        return sph.scope_text.strip()
    if sph and sph.scope_items:
        lines = []
        for item in sph.scope_items:
            if isinstance(item, dict):
                text = str(item.get("text", "")).strip()
                if text:
                    mod = str(item.get("module", "")).strip()
                    lines.append(f"• {text}" + (f" ({mod})" if mod else ""))
        if lines:
            return "\n".join(lines)
    return "—"


def _resolve_po_due(project: Project, po: ProjectPo | None) -> date | None:
    if po and po.po_due_date:
        return po.po_due_date
    return project.po_due_date or project.planned_end_date


def _project_duration_label(db: Session, project_id: int, sph: ProjectSph | None) -> str:
    rows = timeline_display_rows_for_project(db, project_id)
    total = 0
    for r in rows:
        dur = r.get("display_duration_days") or r.get("duration_days")
        if dur and int(dur) > 0:
            total += int(dur)
    if total > 0:
        label = f"{total} hari kerja (total timeline)"
        if sph and sph.planned_md:
            label += f" · {sph.planned_md:g} MD (rencana SPH)"
        return label
    if sph and sph.planned_md:
        return f"{sph.planned_md:g} mandays (rencana SPH)"
    start, end = schedule_bounds(db, project_id)
    if start and end:
        return f"{start.isoformat()} s.d. {end.isoformat()}"
    return "—"


def _period_label(report_date: date, period_start: date, week_end: date) -> str:
    return f"{_fmt_date_slash(period_start)} s.d. {_fmt_date_slash(week_end)}"


def _rag_overall_color(health: dict | None) -> RGBColor:
    if not health:
        return COLOR_MUTED
    key = str(health.get("rag_overall") or "").lower()
    if key == "green":
        return COLOR_GREEN
    if key in ("yellow", "amber"):
        return RGBColor(0xCA, 0x8A, 0x04)
    if key == "red":
        return RGBColor(0xDC, 0x26, 0x26)
    return COLOR_MUTED


def _add_rag_overall_kpi(slide, left, top, health: dict | None) -> None:
    """KPI RAG: lingkaran warna overall project (tanpa teks gap/schedule)."""
    w = Inches(2.45)
    hdr_h = Inches(0.32)
    hdr_shape = slide.shapes.add_shape(1, left, top, w, hdr_h)
    hdr_shape.fill.solid()
    hdr_shape.fill.fore_color.rgb = COLOR_NAVY
    hdr_shape.line.fill.background()
    _add_textbox(
        slide,
        left,
        top + Inches(0.02),
        w,
        hdr_h,
        "RAG Overall",
        size=KPI_LABEL_SIZE,
        bold=True,
        color=COLOR_WHITE,
        align=PP_ALIGN.CENTER,
    )
    body_top = top + hdr_h
    body = slide.shapes.add_shape(1, left, body_top, w, Inches(0.58))
    body.fill.solid()
    body.fill.fore_color.rgb = RGBColor(0xF8, 0xFA, 0xFC)
    body.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
    circle = Inches(0.42)
    cx = left + (w - circle) / 2
    cy = body_top + Inches(0.08)
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx, cy, circle, circle)
    dot.fill.solid()
    dot.fill.fore_color.rgb = _rag_overall_color(health)
    dot.line.color.rgb = RGBColor(0xCF, 0xD8, 0xE3)
    dot.line.width = Pt(1.25)


def _format_pptx_highlights(insights: dict, *, extra_notes: str = "") -> str:
    lines: list[str] = []
    phase_gaps = insights.get("phase_gaps") or []
    lines.append("GAP timeline phase (ketinggalan vs target):")
    if phase_gaps:
        for g in phase_gaps[:6]:
            lines.append(
                f"• {g['name']}: target {g['planned_pct']:.1f}% vs actual "
                f"{g['actual_pct']:.1f}% (selisih {g['gap_pp']:.1f} p.p.)"
            )
    else:
        lines.append("• Tidak ada phase dengan ketinggalan signifikan vs target timeline.")
    lines.append("")
    lines.append("Task selesai minggu ini:")
    tasks_done = insights.get("tasks_completed") or []
    if tasks_done:
        for t in tasks_done[:10]:
            due = t.get("due_date") or "—"
            lines.append(f"• {t['name']} ({t.get('status') or 'done'}, due {due})")
    elif insights.get("phases_current_week"):
        for p in (insights.get("phases_current_week") or [])[:6]:
            lines.append(
                f"• {p['name']}: target {p['planned_pct']:.1f}% · actual {p['actual_pct']:.1f}%"
            )
    else:
        lines.append("• (Belum terdeteksi task/fase di periode ini.)")
    lines.append("")
    next_start = insights.get("next_period_start") or "?"
    next_end = insights.get("next_period_end") or "?"
    lines.append(f"Rencana minggu depan ({next_start} s/d {next_end}):")
    tasks_next = insights.get("tasks_next_week") or []
    if tasks_next:
        for t in tasks_next[:10]:
            due = t.get("due_date") or "—"
            lines.append(f"• {t['name']} ({t.get('status') or 'open'}, due {due})")
    elif insights.get("phases_next_week"):
        for p in (insights.get("phases_next_week") or [])[:6]:
            sd = p.get("start_date") or "—"
            td = p.get("target_date") or "—"
            lines.append(f"• {p['name']} ({sd} s/d {td})")
    else:
        lines.append("• (Belum terdeteksi task/fase terjadwal.)")
    notes = (extra_notes or "").strip()
    if notes:
        lines.extend(["", "Catatan tambahan:", notes])
    return "\n".join(lines)


def _compute_auto_highlights(
    db: Session,
    project: Project,
    report: WeeklyReport,
    *,
    planned: float,
    actual: float,
    period_start: date,
    week_end: date,
) -> str:
    from app.services.schedule import kickoff_milestones, planned_progress_rows
    from app.services.weekly_report_insights import weekly_report_preview_insights

    frozen = report.frozen_metrics or {}
    health = frozen.get("health") or {}
    plan_rows = kickoff_milestones(db, project.id)
    if not plan_rows:
        plan_rows = planned_progress_rows(db, project.id)
    milestones = list(
        db.scalars(select(Milestone).where(Milestone.project_id == project.id)).all()
    )
    tasks = list(
        db.scalars(
            select(ClickUpTaskCache).where(ClickUpTaskCache.project_id == project.id)
        ).all()
    )
    insights = weekly_report_preview_insights(
        db,
        project,
        anchor=report.week_start,
        period_start=period_start,
        cut_off=week_end,
        plan_rows=plan_rows,
        planned_pct=float(planned),
        actual_pct=float(actual),
        milestones=milestones,
        tasks=tasks,
        rag_gap=health.get("rag_gap"),
        rag_schedule=health.get("rag_schedule"),
    )
    summary = report.summary or {}
    manual = str(summary.get("highlights") or "").strip()
    return _format_pptx_highlights(insights, extra_notes=manual)


def _status_display(row: dict) -> str:
    cu = row.get("clickup_status")
    if cu:
        mapping = {
            "done": "Done",
            "in_progress": "In Progress",
            "not_started": "Not Started",
            "open": "Not Started",
        }
        return mapping.get(str(cu).lower(), str(cu).replace("_", " ").title())
    st = str(row.get("status") or "open").lower()
    if st in ("done", "closed"):
        return "Done"
    if st == "in_progress":
        return "In Progress"
    return "Not Started"


def _set_run_font(run, *, size: int, bold: bool = False, color: RGBColor | None = None):
    run.font.name = FONT_FACE
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def _compact_cell_margins(cell) -> None:
    cell.margin_top = Pt(TABLE_CELL_MARGIN_PT)
    cell.margin_bottom = Pt(TABLE_CELL_MARGIN_PT)
    cell.margin_left = Pt(4)
    cell.margin_right = Pt(4)
    tf = cell.text_frame
    tf.word_wrap = True
    for p in tf.paragraphs:
        p.space_before = Pt(0)
        p.space_after = Pt(0)


def _style_table_row(table, row_idx: int, *, header: bool = False) -> None:
    table.rows[row_idx].height = Pt(TABLE_HEADER_ROW_PT if header else TABLE_DATA_ROW_PT)
    for cell in table.rows[row_idx].cells:
        _compact_cell_margins(cell)


def _set_cell_text(
    cell,
    text: str,
    *,
    size: int = 9,
    bold: bool = False,
    color: RGBColor = COLOR_DARK,
    align=PP_ALIGN.LEFT,
) -> None:
    cell.text = text
    _compact_cell_margins(cell)
    for p in cell.text_frame.paragraphs:
        p.alignment = align
        for run in p.runs:
            _set_run_font(run, size=size, bold=bold, color=color)


def _style_chart(chart) -> None:
    chart.has_title = True
    if chart.chart_title and chart.chart_title.text_frame:
        chart.chart_title.text_frame.text = chart.chart_title.text_frame.text[:60]
        for p in chart.chart_title.text_frame.paragraphs:
            for run in p.runs:
                _set_run_font(run, size=11, bold=True, color=COLOR_NAVY)
    try:
        chart.category_axis.tick_labels.font.name = FONT_FACE
        chart.category_axis.tick_labels.font.size = Pt(9)
        chart.value_axis.tick_labels.font.name = FONT_FACE
        chart.value_axis.tick_labels.font.size = Pt(9)
        if chart.legend:
            chart.legend.font.name = FONT_FACE
            chart.legend.font.size = Pt(9)
    except AttributeError:
        pass


def _add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    *,
    size: int = 14,
    bold: bool = False,
    color: RGBColor = COLOR_DARK,
    align=PP_ALIGN.LEFT,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run_font(run, size=size, bold=bold, color=color)
    return box


def _slide_body_backdrop(slide) -> None:
    """Area konten di bawah judul — konsisten slide 2–4."""
    backdrop = slide.shapes.add_shape(
        1, Inches(0.35), Inches(0.92), Inches(12.65), Inches(6.05)
    )
    backdrop.fill.solid()
    backdrop.fill.fore_color.rgb = RGBColor(0xFA, 0xFB, 0xFC)
    backdrop.line.color.rgb = RGBColor(0xE8, 0xED, 0xF2)


def _slide_footer(slide, page: int):
    bar = slide.shapes.add_shape(
        1, Inches(0), Inches(7.05), Inches(4.2), Inches(0.35)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_ORANGE
    bar.line.fill.background()
    _add_textbox(
        slide,
        Inches(0.15),
        Inches(7.08),
        Inches(3),
        Inches(0.3),
        "www.ebesha.net",
        size=BODY_SIZE,
        bold=True,
        color=COLOR_WHITE,
    )
    _add_textbox(
        slide,
        Inches(12.5),
        Inches(7.08),
        Inches(0.6),
        Inches(0.3),
        str(page),
        size=BODY_SIZE,
        bold=True,
        color=COLOR_NAVY,
        align=PP_ALIGN.RIGHT,
    )


def _slide_header(slide, title: str):
    _add_textbox(
        slide,
        Inches(0.45),
        Inches(0.35),
        Inches(8),
        Inches(0.55),
        title,
        size=TITLE_SIZE,
        bold=True,
        color=COLOR_NAVY,
    )
    accent = slide.shapes.add_shape(
        1, Inches(11.2), Inches(0.15), Inches(1.8), Inches(0.55)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = COLOR_ORANGE
    accent.line.fill.background()
    accent2 = slide.shapes.add_shape(
        1, Inches(11.55), Inches(0.35), Inches(1.5), Inches(0.45)
    )
    accent2.fill.solid()
    accent2.fill.fore_color.rgb = COLOR_NAVY
    accent2.line.fill.background()


def _blank_slide(prs: Presentation):
    layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
    return prs.slides.add_slide(layout)


def _build_slide_title(prs: Presentation, project: Project, report_date: date):
    slide = _blank_slide(prs)
    bg = _title_bg_path()
    if bg:
        slide.shapes.add_picture(str(bg), Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    else:
        panel = slide.shapes.add_shape(
            1, Inches(0), Inches(0), Inches(8.8), prs.slide_height
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = COLOR_ORANGE
        panel.line.fill.background()
        overlay = slide.shapes.add_shape(
            1, Inches(0.55), Inches(2.35), Inches(7.2), Inches(2.35)
        )
        overlay.fill.solid()
        overlay.fill.fore_color.rgb = RGBColor(0x33, 0x33, 0x33)
        overlay.fill.transparency = 0.35
        overlay.line.fill.background()

    # Teks di area orange kiri (sesuai template cover)
    text_left = Inches(0.85)
    text_width = Inches(7.4)
    _add_textbox(
        slide,
        text_left,
        Inches(2.55),
        text_width,
        Inches(0.6),
        "Weekly Report Project",
        size=TITLE_SIZE,
        bold=True,
        color=COLOR_WHITE,
    )
    name_lines = project.name.strip()
    if len(name_lines) > 52:
        name_lines = name_lines[:49] + "…"
    _add_textbox(
        slide,
        text_left,
        Inches(3.25),
        text_width,
        Inches(1.05),
        name_lines,
        size=TITLE_SIZE,
        bold=True,
        color=COLOR_WHITE,
    )
    date_str = _fmt_date_en_long(report_date)
    _add_textbox(
        slide,
        text_left,
        Inches(4.45),
        text_width,
        Inches(0.45),
        date_str,
        size=14,
        color=COLOR_WHITE,
    )
    underline = slide.shapes.add_shape(
        1, text_left, Inches(4.92), Inches(2.8), Inches(0.015)
    )
    underline.fill.solid()
    underline.fill.fore_color.rgb = COLOR_WHITE
    underline.line.fill.background()

    pm = (project.project_manager or "LMD").strip()
    _add_textbox(
        slide,
        Inches(0.55),
        Inches(6.55),
        Inches(4),
        Inches(0.35),
        f"Prepared By: {pm}",
        size=11,
        color=COLOR_DARK,
    )


def _build_slide_overview(
    prs: Presentation,
    project: Project,
    sph: ProjectSph | None,
    po: ProjectPo | None,
    db: Session,
):
    slide = _blank_slide(prs)
    _slide_header(slide, "Overview Project")
    _slide_body_backdrop(slide)
    _slide_footer(slide, 2)

    brief = _resolve_brief(project, sph)
    duration = _project_duration_label(db, project.id, sph)
    target = _resolve_po_due(project, po)
    pm = project.project_manager or "—"

    rows = [
        ("Project Name", project.name),
        ("Project Brief", brief),
        ("Project Duration", duration),
        ("Project Target Date", _fmt_date_en_long(target) if target else "—"),
        ("Project Manager", pm),
    ]

    n_rows = len(rows) + 1
    table_h = Inches((TABLE_HEADER_ROW_PT * 1.2 + TABLE_DATA_ROW_PT * len(rows)) / 72 + 0.15)
    table_shape = slide.shapes.add_table(
        n_rows, 2, Inches(0.55), Inches(1.08), Inches(12.2), table_h
    )
    table = table_shape.table
    table.columns[0].width = Inches(2.35)
    table.columns[1].width = Inches(9.85)

    hdr = table.cell(0, 0)
    hdr.merge(table.cell(0, 1))
    _set_cell_text(
        hdr,
        project.name,
        size=TABLE_HEADER_FONT,
        bold=True,
        color=COLOR_WHITE,
        align=PP_ALIGN.CENTER,
    )
    hdr.fill.solid()
    hdr.fill.fore_color.rgb = COLOR_ORANGE
    _style_table_row(table, 0, header=True)

    for i, (label, value) in enumerate(rows, start=1):
        lc = table.cell(i, 0)
        vc = table.cell(i, 1)
        brief = label == "Project Brief"
        _set_cell_text(
            lc,
            label,
            size=TABLE_BODY_FONT,
            bold=True,
            color=COLOR_DARK,
            align=PP_ALIGN.RIGHT,
        )
        _set_cell_text(
            vc,
            value,
            size=TABLE_BODY_FONT,
            bold=False,
            color=COLOR_DARK,
            align=PP_ALIGN.LEFT,
        )
        lc.fill.solid()
        vc.fill.solid()
        fill = COLOR_PEACH if i % 2 else COLOR_PEACH_LIGHT
        lc.fill.fore_color.rgb = fill
        vc.fill.fore_color.rgb = COLOR_WHITE
        _style_table_row(table, i, header=False)


def _scurve_snapshot_rows(
    points: list[dict],
    report_date: date,
    report_cut_off: date,
    *,
    max_rows: int = 8,
) -> list[dict]:
    """Rows up to report week; actual kosong untuk tanggal > cut-off laporan."""
    rows: list[dict] = []
    for pt in points:
        d = _point_anchor_date(pt)
        if not d or d > report_cut_off:
            continue
        masked = {
            **pt,
            "_date": d,
            "actual_pct": _actual_pct_for_report_export(pt, report_cut_off),
        }
        rows.append(masked)
        if d >= report_date:
            break
    if not rows:
        return []
    return rows[-max_rows:]


def _add_scurve_snapshot_table(
    slide,
    snapshot: list[dict],
    *,
    left,
    top,
    width,
    title: str = "S-curve snapshot (weekly report)",
) -> None:
    if not snapshot:
        return
    _add_textbox(
        slide,
        left,
        top,
        width,
        Inches(0.22),
        title,
        size=TABLE_BODY_FONT,
        bold=True,
        color=COLOR_NAVY,
    )
    top = top + Inches(0.24)
    n = len(snapshot) + 1
    table_h = Inches((TABLE_HEADER_ROW_PT + TABLE_DATA_ROW_PT * len(snapshot)) / 72 + 0.12)
    shape = slide.shapes.add_table(n, 4, left, top, width, table_h)
    table = shape.table
    cols = ["Tanggal laporan", "Planned %", "Actual %", "SPI"]
    col_w = [width / 4] * 4
    for i, w in enumerate(col_w):
        table.columns[i].width = int(w)
    for j, h in enumerate(cols):
        cell = table.cell(0, j)
        _set_cell_text(
            cell,
            h,
            size=TABLE_HEADER_FONT,
            bold=True,
            color=COLOR_WHITE,
            align=PP_ALIGN.CENTER,
        )
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_NAVY
    _style_table_row(table, 0, header=True)
    for i, pt in enumerate(snapshot, start=1):
        d = pt["_date"]
        actual = pt.get("actual_pct")
        spi = pt.get("spi")
        spi_txt = (
            f"{float(spi):.4f}" if spi is not None and actual is not None else "—"
        )
        actual_txt = f"{float(actual):.2f}" if actual is not None else "—"
        vals = [
            _fmt_date_slash(d),
            f"{float(pt.get('planned_pct') or 0):.2f}",
            actual_txt,
            spi_txt,
        ]
        fill = COLOR_PEACH_LIGHT if i % 2 else COLOR_WHITE
        for j, val in enumerate(vals):
            cell = table.cell(i, j)
            _set_cell_text(
                cell,
                val,
                size=TABLE_BODY_FONT,
                align=PP_ALIGN.CENTER,
            )
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
        _style_table_row(table, i, header=False)


def _build_slide_progress(
    prs: Presentation,
    project: Project,
    report: WeeklyReport,
    db: Session,
    *,
    planned: float,
    actual: float,
    highlights: str,
    period_start: date,
    week_end: date,
):
    slide = _blank_slide(prs)
    _slide_header(slide, "Progress Milestone")
    _slide_body_backdrop(slide)
    _slide_footer(slide, 3)

    period = _period_label(report.week_start, period_start, week_end)
    frozen = report.frozen_metrics or {}
    health = frozen.get("health") if isinstance(frozen.get("health"), dict) else {}
    spi_val = frozen.get("spi")
    spi_txt = f"{float(spi_val):.4f}" if spi_val is not None else "—"

    panel = slide.shapes.add_shape(
        1, Inches(0.45), Inches(0.98), Inches(5.55), Inches(6.0)
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = COLOR_WHITE
    panel.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)

    band = slide.shapes.add_shape(
        1, Inches(0.45), Inches(0.98), Inches(5.55), Inches(0.34)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = COLOR_NAVY
    band.line.fill.background()
    _add_textbox(
        slide,
        Inches(0.55),
        Inches(1.02),
        Inches(5.35),
        Inches(0.28),
        f"Project Status this Period ({period})",
        size=BODY_SIZE,
        bold=True,
        color=COLOR_WHITE,
    )

    def metric_box(
        left,
        top,
        title: str,
        value: str,
        header_color: RGBColor,
        *,
        value_size: int = KPI_VALUE_SIZE,
    ):
        w = Inches(2.45)
        hdr_h = Inches(0.32)
        hdr_shape = slide.shapes.add_shape(1, left, top, w, hdr_h)
        hdr_shape.fill.solid()
        hdr_shape.fill.fore_color.rgb = header_color
        hdr_shape.line.fill.background()
        _add_textbox(
            slide,
            left,
            top + Inches(0.02),
            w,
            hdr_h,
            title,
            size=KPI_LABEL_SIZE,
            bold=True,
            color=COLOR_WHITE,
            align=PP_ALIGN.CENTER,
        )
        body_top = top + hdr_h
        body = slide.shapes.add_shape(1, left, body_top, w, Inches(0.58))
        body.fill.solid()
        body.fill.fore_color.rgb = RGBColor(0xF8, 0xFA, 0xFC)
        body.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
        _add_textbox(
            slide,
            left,
            body_top + Inches(0.08),
            w,
            Inches(0.48),
            value,
            size=value_size,
            bold=True,
            color=COLOR_DARK,
            align=PP_ALIGN.CENTER,
        )

    row1 = Inches(1.38)
    row2 = Inches(2.38)
    metric_box(Inches(0.55), row1, "Target Progress", f"{planned:.0f}%", COLOR_ORANGE)
    metric_box(Inches(3.15), row1, "Actual Progress", f"{actual:.0f}%", COLOR_GREEN)
    metric_box(Inches(0.55), row2, "SPI", spi_txt, COLOR_NAVY, value_size=16)
    _add_rag_overall_kpi(slide, Inches(3.15), row2, health)

    hl_box = slide.shapes.add_shape(
        1, Inches(0.55), Inches(3.08), Inches(5.05), Inches(3.75)
    )
    hl_box.fill.solid()
    hl_box.fill.fore_color.rgb = COLOR_WHITE
    hl_box.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
    _add_textbox(
        slide,
        Inches(0.65),
        Inches(3.14),
        Inches(4.9),
        Inches(0.22),
        "Highlight",
        size=KPI_LABEL_SIZE,
        bold=True,
        color=COLOR_NAVY,
    )
    hl = (highlights or "—").strip()
    if len(hl) > 850:
        hl = hl[:847] + "…"
    _add_textbox(
        slide,
        Inches(0.65),
        Inches(3.38),
        Inches(4.85),
        Inches(3.35),
        hl,
        size=TABLE_BODY_FONT,
        color=COLOR_DARK,
    )

    points = scurve_points(db, project.id)
    snapshot = _scurve_snapshot_rows(points, report.week_start, week_end)
    chart_top = Inches(1.02)
    chart_h = Inches(3.55) if snapshot else Inches(5.45)

    if points:
        cats = []
        planned_series = []
        actual_series = []
        chart_pts = points if len(points) <= 20 else points[:: max(1, len(points) // 18)]
        for pt in chart_pts:
            rd = pt.get("date") or pt.get("anchor_date") or pt.get("week_start")
            d = _parse_iso(rd if isinstance(rd, str) else None)
            cats.append(_fmt_date_slash(d) if d else str(rd or ""))
            planned_series.append(float(pt.get("planned_pct") or 0))
            actual_series.append(_actual_pct_for_report_export(pt, week_end))

        chart_data = ChartData()
        chart_data.categories = cats
        chart_data.add_series("Target Linear (%)", planned_series)
        chart_data.add_series("Cummulative Actual (%)", actual_series)

        chart_frame = slide.shapes.add_chart(
            XL_CHART_TYPE.LINE_MARKERS,
            Inches(6.15),
            chart_top,
            Inches(6.75),
            chart_h,
            chart_data,
        )
        chart = chart_frame.chart
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = True
        chart.chart_title.text_frame.text = project.name[:55]
        _style_chart(chart)
        if chart.series:
            chart.series[0].format.line.color.rgb = COLOR_ORANGE
            if len(chart.series) > 1:
                chart.series[1].format.line.color.rgb = COLOR_GREEN
        if snapshot:
            _add_scurve_snapshot_table(
                slide,
                snapshot,
                left=Inches(6.15),
                top=Inches(4.72),
                width=Inches(6.75),
            )
    else:
        _add_textbox(
            slide,
            Inches(6.2),
            Inches(2.2),
            Inches(6.5),
            Inches(0.8),
            "S-curve belum tersedia — konfirmasi timeline kick off dan generate target laporan.",
            size=9,
            color=COLOR_MUTED,
        )


def _activity_name(row: dict) -> str:
    depth = int(row.get("depth") or 0)
    prefix = "  " * depth
    return prefix + str(row.get("name") or "—")


def _build_slide_timeline(
    prs: Presentation,
    project: Project,
    db: Session,
    *,
    as_of: date | None = None,
):
    slide = _blank_slide(prs)
    title = f"Progress {project.name[:48]} – Timeline"
    _slide_header(slide, title)
    _slide_body_backdrop(slide)
    _slide_footer(slide, 4)

    rows = timeline_display_rows_for_project(db, project.id)
    display_rows = [r for r in rows if include_in_report_timeline_list(r)]
    max_timeline_rows = 20
    display_rows = display_rows[:max_timeline_rows]

    n = len(display_rows) + 1
    table_h = Inches((TABLE_HEADER_ROW_PT + TABLE_DATA_ROW_PT * len(display_rows)) / 72 + 0.15)
    table_shape = slide.shapes.add_table(
        n, 7, Inches(0.35), Inches(1.02), Inches(12.55), min(table_h, Inches(5.82))
    )
    table = table_shape.table
    headers = [
        "No",
        "Activity",
        "Start Date",
        "Due Date",
        "Progress",
        "Status",
        "Keterangan",
    ]
    widths = [0.42, 3.55, 1.1, 1.1, 0.78, 0.98, 2.35]
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)

    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        _set_cell_text(
            cell,
            h,
            size=TABLE_HEADER_FONT,
            bold=True,
            color=COLOR_WHITE,
            align=PP_ALIGN.CENTER,
        )
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_ORANGE
    _style_table_row(table, 0, header=True)

    cut_off = as_of or today_jakarta()
    ms_list = kickoff_milestones(db, project.id) or list(
        db.scalars(select(Milestone).where(Milestone.project_id == project.id)).all()
    )
    caches = list(
        db.scalars(
            select(ClickUpTaskCache).where(ClickUpTaskCache.project_id == project.id)
        ).all()
    )
    ms_by_id = {m.id: m for m in ms_list}

    with clickup_status_mapping_context(db):
        tasks_by_id, milestone_cache = build_clickup_lookups(ms_list, caches)
        progress_ctx = ReportRowProgressContext(
            db=db,
            project=project,
            milestones=ms_list,
            ms_by_id=ms_by_id,
            tasks_by_id=tasks_by_id,
            milestone_cache=milestone_cache,
            caches=caches,
        )
        for i, row in enumerate(display_rows, start=1):
            start_d = _parse_iso(row.get("display_start") or row.get("start_date"))
            end_d = _parse_iso(row.get("display_end") or row.get("target_date"))
            pct = display_row_actual_pct(progress_ctx, row, cut_off)
            pct_txt = f"{pct:.0f}%"
            notes = (row.get("notes") or "").strip()
            if not notes and row.get("module"):
                notes = str(row.get("module"))
            if len(notes) > 120:
                notes = notes[:117] + "…"
            values = [
                str(i),
                _activity_name(row),
                _fmt_date_slash(start_d),
                _fmt_date_slash(end_d),
                pct_txt,
                _status_display(row),
                notes or "—",
            ]
            fill = COLOR_PEACH if i % 2 == 0 else COLOR_PEACH_LIGHT
            for j, val in enumerate(values):
                cell = table.cell(i, j)
                align = PP_ALIGN.CENTER if j in (0, 2, 3, 4, 5) else PP_ALIGN.LEFT
                bold = j == 1 and int(row.get("depth") or 0) == 0
                _set_cell_text(
                    cell,
                    val,
                    size=TABLE_BODY_FONT,
                    bold=bold,
                    color=COLOR_DARK,
                    align=align,
                )
                cell.fill.solid()
                cell.fill.fore_color.rgb = fill if j != 0 else COLOR_PEACH_LIGHT
            _style_table_row(table, i, header=False)


def build_weekly_report_pptx(
    dest: Path,
    *,
    db: Session,
    project: Project,
    report: WeeklyReport,
    sph: ProjectSph | None = None,
    po: ProjectPo | None = None,
) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    frozen = report.frozen_metrics or {}
    planned = float(frozen.get("planned_pct") or 0)
    actual = float(frozen.get("actual_pct") or 0)

    pstart = project_report_start_date(db, project)
    _, period_start, week_end = resolve_report_period(
        project.weekly_report_anchor_weekday,
        project.weekly_report_cutoff_offset_days,
        report.week_start,
        project_start_date=pstart,
        explicit_first_report_date=project.weekly_report_first_anchor_date,
    )
    try:
        highlights = _compute_auto_highlights(
            db,
            project,
            report,
            planned=planned,
            actual=actual,
            period_start=period_start,
            week_end=week_end,
        )
    except Exception:
        summary = report.summary or {}
        highlights = str(summary.get("highlights") or "—")

    _build_slide_title(prs, project, report.week_start)
    _build_slide_overview(prs, project, sph, po, db)
    _build_slide_progress(
        prs,
        project,
        report,
        db,
        planned=planned,
        actual=actual,
        highlights=highlights,
        period_start=period_start,
        week_end=week_end,
    )
    _build_slide_timeline(prs, project, db, as_of=week_end)

    prs.save(str(dest))
