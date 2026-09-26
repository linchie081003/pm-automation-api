"""Parse holiday dates from Excel uploads."""
from datetime import date, datetime, timedelta
from io import BytesIO

from openpyxl import load_workbook


def _cell_to_date(val) -> date | None:
    if val is None:
        return None
    if isinstance(val, date) and not isinstance(val, datetime):
        return val
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, (int, float)):
        base = date(1899, 12, 30)
        return base + timedelta(days=int(val))
    s = str(val).strip()[:10]
    if len(s) >= 10:
        try:
            return date.fromisoformat(s)
        except ValueError:
            return None
    return None


def parse_holidays_xlsx(content: bytes) -> list[dict]:
    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header = [str(c or "").strip().lower() for c in rows[0]]
    date_col = None
    label_col = None
    for i, h in enumerate(header):
        if h in ("tanggal", "date", "tgl", "holiday"):
            date_col = i
        if h in ("keterangan", "label", "nama", "description"):
            label_col = i
    start_row = 1 if date_col is not None or label_col is not None else 0
    if date_col is None:
        date_col = 0

    out: list[dict] = []
    seen: set[str] = set()
    for row in rows[start_row:]:
        if not row or date_col >= len(row):
            continue
        d = _cell_to_date(row[date_col])
        if not d:
            continue
        key = d.isoformat()
        if key in seen:
            continue
        seen.add(key)
        label = ""
        if label_col is not None and label_col < len(row) and row[label_col]:
            label = str(row[label_col]).strip()
        out.append({"date": key, "label": label})
    wb.close()
    return out
