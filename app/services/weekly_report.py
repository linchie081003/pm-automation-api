import uuid
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import settings
from app.models import (
    ClickUpTaskCache,
    Milestone,
    ProgressSnapshotSource,
    Project,
    ProjectHealthSnapshot,
    WeeklyReport,
)
from app.services.health import compute_health
from app.services.schedule import (
    baseline_for_date,
    compute_spi,
    planned_pct_as_of,
    save_weekly_progress,
)
from app.services.schedule_window import kickoff_milestones, project_report_start_date
from app.services.progress import resolve_actual_progress
from app.services.report_calendar import (
    active_open_report_date,
    ensure_weekly_period_has_started,
    resolve_report_period,
)
from app.services.progress_metrics import snapshot_for_anchor_week
from app.services.clickup import sync_project_tasks
from app.services.templates.loader import copy_template
from app.services.templates.placeholders import build_mapping, replace_in_pptx, replace_in_xlsx
from app.services.weekly_report_pptx import build_weekly_report_pptx
from app.services.yyyymmdd_report_excel import export_yyyymmdd_workbook


def _ensure_clickup_ready(db: Session, project: Project) -> None:
    if not project.clickup_enabled:
        return
    if project.clickup_provision_status != "provisioned":
        raise ValueError("Generate task ClickUp dari timeline terlebih dahulu (tab SETTING).")
    sync_project_tasks(db, project)
    db.flush()
    has_task = db.scalar(
        select(ClickUpTaskCache.id)
        .where(ClickUpTaskCache.project_id == project.id)
        .limit(1)
    )
    if not has_task:
        raise ValueError("Belum ada task ClickUp — sync gagal atau list kosong.")


def _next_weekly_report_doc_version(
    db: Session, project_id: int, stamp: str, code: str
) -> int:
    from app.models import Document

    filenames = db.scalars(
        select(Document.filename).where(
            Document.project_id == project_id,
            Document.filename.like(f"{stamp}_Weekly_Report_{code}%"),
        )
    ).all()
    max_v = 0
    for fn in filenames:
        if "_v" not in fn:
            max_v = max(max_v, 1)
            continue
        try:
            tail = fn.rsplit("_v", 1)[-1]
            max_v = max(max_v, int(tail.split(".")[0]))
        except ValueError:
            continue
    return max_v + 1


def _report_file_basename(stamp: str, code: str, version: int) -> str:
    return f"{stamp}_Weekly_Report_{code}_v{version:02d}"


def _resolve_preview_metrics(
    db: Session,
    project: Project,
    report_date: date,
    week_end: date,
    existing: WeeklyReport | None,
) -> tuple[float, float, float | None, dict, str, bool]:
    """
    Returns planned, actual, spi, health, metrics_source, matches_project_health.
    """
    snap = snapshot_for_anchor_week(db, project.id, report_date)
    active_rd = active_open_report_date(date.today(), project.weekly_report_anchor_weekday)
    is_active_period = report_date == active_rd

    if existing and existing.frozen_metrics:
        fm = existing.frozen_metrics
        planned = float(fm.get("planned_pct", 0))
        actual = float(fm.get("actual_pct", 0))
        spi = fm.get("spi")
        spi_f = float(spi) if spi is not None else compute_spi(actual, planned)
        health = dict(fm.get("health") or {})
        if not health:
            health = compute_health(db, project.id, week_end)
        return planned, actual, spi_f, health, "saved_report", False

    frozen_sources = {
        ProgressSnapshotSource.weekly_report,
        ProgressSnapshotSource.manual_save,
    }
    if snap and snap.source in frozen_sources:
        planned = float(snap.planned_cumulative_pct)
        actual = float(snap.actual_cumulative_pct)
        spi_f = float(snap.spi_at_week)
        health = compute_health(db, project.id, week_end)
        return planned, actual, spi_f, health, "snapshot", False

    if is_active_period:
        health = compute_health(db, project.id)
        planned = float(health.get("planned_progress_pct") or 0)
        actual = float(health.get("actual_progress_pct") or 0)
        spi_f = health.get("spi")
        if spi_f is None:
            spi_f = compute_spi(actual, planned)
        else:
            spi_f = float(spi_f)
        return planned, actual, spi_f, health, "active_live", True

    plan_rows = kickoff_milestones(db, project.id)
    if not plan_rows:
        from app.services.schedule import planned_progress_rows

        plan_rows = planned_progress_rows(db, project.id)
    planned = planned_pct_as_of(plan_rows, week_end, db=db)
    actual = resolve_actual_progress(db, project, week_end)
    spi_f = compute_spi(actual, planned)
    health = compute_health(db, project.id, week_end)
    return float(planned), float(actual), spi_f, health, "computed_historical", False


def preview_weekly_report(
    db: Session,
    project_id: int,
    week_start: date,
    notes: str | None = None,
) -> dict:
    project = db.get(Project, project_id)
    if not project:
        raise ValueError("Project not found")
    week_start = ensure_weekly_period_has_started(
        project.weekly_report_anchor_weekday,
        project.weekly_report_cutoff_offset_days,
        week_start,
    )
    pstart = project_report_start_date(db, project_id)
    report_date, period_start, week_end = resolve_report_period(
        project.weekly_report_anchor_weekday,
        project.weekly_report_cutoff_offset_days,
        week_start,
        project_start_date=pstart,
        explicit_first_report_date=project.weekly_report_first_anchor_date,
    )
    week_start = report_date
    existing = db.scalar(
        select(WeeklyReport).where(
            WeeklyReport.project_id == project_id,
            WeeklyReport.week_start == week_start,
        )
    )

    planned, actual, spi, health, metrics_source, matches_project_health = (
        _resolve_preview_metrics(db, project, report_date, week_end, existing)
    )
    baseline = baseline_for_date(db, project_id, week_end)
    plan_rows = kickoff_milestones(db, project_id)
    if not plan_rows:
        from app.services.schedule import planned_progress_rows

        plan_rows = planned_progress_rows(db, project_id)

    milestones = db.scalars(
        select(Milestone).where(Milestone.project_id == project_id)
    ).all()
    tasks = db.scalars(
        select(ClickUpTaskCache).where(ClickUpTaskCache.project_id == project_id)
    ).all()

    from app.services.weekly_report_insights import weekly_report_preview_insights

    insights = weekly_report_preview_insights(
        db,
        project,
        anchor=week_start,
        period_start=period_start,
        cut_off=week_end,
        plan_rows=plan_rows,
        planned_pct=float(planned),
        actual_pct=float(actual),
        milestones=list(milestones),
        tasks=list(tasks),
        rag_gap=health.get("rag_gap"),
        rag_schedule=health.get("rag_schedule"),
    )

    from app.services.report_calendar import display_period_day_count

    return {
        "report_date": report_date.isoformat(),
        "week_start": week_start.isoformat(),
        "week_end": week_end.isoformat(),
        "anchor_date": report_date.isoformat(),
        "period_start": period_start.isoformat(),
        "target_week_start": period_start.isoformat(),
        "cut_off_date": week_end.isoformat(),
        "period_length_days": project.weekly_report_cutoff_offset_days,
        "period_day_count": display_period_day_count(project.weekly_report_cutoff_offset_days),
        "status_date_report": health.get("status_date") or week_end.isoformat(),
        "generate_progress_pct": round(actual, 2),
        "metrics_source": metrics_source,
        "matches_project_health": matches_project_health,
        "already_exists": existing is not None,
        "existing_report_id": existing.id if existing else None,
        "next_document_version": _next_weekly_report_doc_version(
            db, project_id, week_start.strftime("%Y%m%d"), project.code
        ),
        "project_code": project.code,
        "project_name": project.name,
        "baseline_version": baseline.version if baseline else None,
        "planned_pct": planned,
        "actual_pct": actual,
        "spi": spi,
        "deviation_pct": insights["deviation_pct"],
        "deviation_pp": insights["deviation_pp"],
        "gap_pp": insights["gap_pp"],
        "rag_gap": insights["rag_gap"],
        "rag_schedule": insights["rag_schedule"],
        "rag_deviation": insights["rag_deviation"],
        "phase_gaps": insights["phase_gaps"],
        "phases_current_week": insights["phases_current_week"],
        "phases_next_week": insights["phases_next_week"],
        "use_phase_fallback": insights["use_phase_fallback"],
        "use_phase_next_fallback": insights["use_phase_next_fallback"],
        "tasks_completed": insights["tasks_completed"],
        "tasks_next_week": insights["tasks_next_week"],
        "next_period_start": insights["next_period_start"],
        "next_period_end": insights["next_period_end"],
        "highlights_draft": insights["highlights_draft"],
        "rag_overall": health.get("rag_overall"),
        "health": health,
        "milestone_count": len(milestones),
        "task_count": len(tasks),
        "milestones_preview": [
            {"name": m.name, "status": m.status.value, "target_date": str(m.target_date or "")}
            for m in milestones[:10]
        ],
        "tasks_preview": [
            {
                "name": t.name,
                "status": t.status,
                "due_date": t.due_date.isoformat() if t.due_date else None,
            }
            for t in tasks[:10]
        ],
        "output_formats": ["xlsx", "pptx"],
        "notes": notes or insights["highlights_draft"] or project.weekly_notes or "",
    }


def compose_weekly_report_notes(notes: str | None, mitigation_plan: str | None) -> str:
    parts: list[str] = []
    if notes and notes.strip():
        parts.append(notes.strip())
    if mitigation_plan and mitigation_plan.strip():
        parts.append("Mitigasi mengejar plan:\n" + mitigation_plan.strip())
    return "\n\n".join(parts)


def generate_weekly_report(
    db: Session,
    project_id: int,
    week_start: date,
    user_id: int,
    notes: str | None,
    mitigation_plan: str | None = None,
    *,
    regenerate: bool = False,
) -> WeeklyReport:
    project = db.get(Project, project_id)
    if not project:
        raise ValueError("Project not found")
    week_start = ensure_weekly_period_has_started(
        project.weekly_report_anchor_weekday,
        project.weekly_report_cutoff_offset_days,
        week_start,
    )
    _ensure_clickup_ready(db, project)
    pstart = project_report_start_date(db, project_id)
    report_date, _, week_end = resolve_report_period(
        project.weekly_report_anchor_weekday,
        project.weekly_report_cutoff_offset_days,
        week_start,
        project_start_date=pstart,
        explicit_first_report_date=project.weekly_report_first_anchor_date,
    )
    week_start = report_date
    existing = db.scalar(
        select(WeeklyReport).where(
            WeeklyReport.project_id == project_id,
            WeeklyReport.week_start == week_start,
        )
    )
    if existing and not regenerate:
        return existing

    active_rd = active_open_report_date(date.today(), project.weekly_report_anchor_weekday)
    if report_date == active_rd:
        snap = save_weekly_progress(
            db, project_id, week_start, ProgressSnapshotSource.weekly_report
        )
    else:
        snap = snapshot_for_anchor_week(db, project_id, week_start)
        if not snap:
            raise ValueError(
                "Snapshot progress untuk periode ini belum ada — tidak dapat generate dokumen."
            )

    planned, actual, spi, health, _, _ = _resolve_preview_metrics(
        db, project, report_date, week_end, None
    )
    if report_date == active_rd:
        health = compute_health(db, project_id)
        planned = float(health.get("planned_progress_pct") or planned)
        actual = float(health.get("actual_progress_pct") or actual)
        spi = health.get("spi") if health.get("spi") is not None else compute_spi(actual, planned)
        snap.planned_cumulative_pct = planned
        snap.actual_cumulative_pct = actual
        snap.spi_at_week = float(spi) if spi is not None else 0.0

    milestones = db.scalars(
        select(Milestone).where(Milestone.project_id == project_id)
    ).all()
    tasks = db.scalars(
        select(ClickUpTaskCache).where(ClickUpTaskCache.project_id == project_id)
    ).all()
    combined_notes = compose_weekly_report_notes(notes, mitigation_plan)

    if existing and regenerate:
        report = existing
        summary = dict(report.summary or {})
        summary.update(
            {
                "highlights": combined_notes or project.weekly_notes or "",
                "mitigation_plan": (mitigation_plan or "").strip(),
                "milestones": [{"name": m.name, "status": m.status.value} for m in milestones],
                "task_count": len(tasks),
            }
        )
        report.summary = summary
    else:
        summary = {
            "highlights": combined_notes or project.weekly_notes or "",
            "mitigation_plan": (mitigation_plan or "").strip(),
            "milestones": [{"name": m.name, "status": m.status.value} for m in milestones],
            "task_count": len(tasks),
        }
        report = WeeklyReport(
            project_id=project_id,
            week_start=week_start,
            week_end=week_end,
            baseline_version=snap.baseline_version,
            summary=summary,
            frozen_metrics={},
            generated_by_id=user_id,
        )
        db.add(report)
        db.flush()

    frozen = {
        "planned_pct": snap.planned_cumulative_pct,
        "actual_pct": snap.actual_cumulative_pct,
        "spi": snap.spi_at_week,
        "health": health,
    }
    report.frozen_metrics = frozen
    report.baseline_version = snap.baseline_version
    report.week_end = week_end
    snap.weekly_report_id = report.id

    stamp = report.week_start.strftime("%Y%m%d")
    doc_version = _next_weekly_report_doc_version(db, project_id, stamp, project.code)
    xlsx_path, pptx_path = _write_report_files(
        db, project, report, milestones, tasks, file_version=doc_version
    )
    report.xlsx_path = xlsx_path
    report.pptx_path = pptx_path
    summary = dict(report.summary or {})
    summary["document_version"] = doc_version
    report.summary = summary

    from app.models import DocumentType
    from app.services.document_registry import register_file_as_document

    base = _report_file_basename(stamp, project.code, doc_version)
    if xlsx_path:
        register_file_as_document(
            db,
            project,
            user_id,
            Path(xlsx_path),
            DocumentType.progress_report,
            display_filename=f"{base}.xlsx",
        )
    if pptx_path:
        register_file_as_document(
            db,
            project,
            user_id,
            Path(pptx_path),
            DocumentType.progress_report,
            display_filename=f"{base}.pptx",
        )

    if not (existing and regenerate):
        db.add(
            ProjectHealthSnapshot(
                project_id=project_id,
                as_of_date=week_end,
                planned_progress_pct=snap.planned_cumulative_pct,
                actual_progress_pct=snap.actual_cumulative_pct,
                spi=snap.spi_at_week,
                rag_schedule=health["rag_schedule"],
                rag_gap=health["rag_gap"],
                rag_overall=health["rag_overall"],
                baseline_version=snap.baseline_version,
                weekly_report_id=report.id,
            )
        )
    db.flush()
    return report


def _task_table(tasks) -> str:
    return "\n".join(
        f"- {t.name} | {t.status} | due {t.due_date or '-'}" for t in tasks[:50]
    )


def _write_report_files(
    db: Session,
    project: Project,
    report: WeeklyReport,
    milestones,
    tasks,
    *,
    file_version: int = 1,
) -> tuple[str, str]:
    out_dir = Path(settings.upload_dir) / "reports" / str(project.id)
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = report.week_start.strftime("%Y%m%d")
    file_base = _report_file_basename(stamp, project.code, file_version)
    summary = report.summary or {}
    highlights = summary.get("highlights") or ""
    mitigation = summary.get("mitigation_plan") or ""
    mapping = build_mapping(
        PROJECT_NAME=project.name,
        PROJECT_CODE=project.code,
        CLIENT_NAME=project.client_name,
        WEEK=str(report.week_start),
        PLANNED_PCT=str(report.frozen_metrics.get("planned_pct")),
        ACTUAL_PCT=str(report.frozen_metrics.get("actual_pct")),
        SPI=str(report.frozen_metrics.get("spi")),
        HIGHLIGHTS=highlights,
        MITIGATION_PLAN=mitigation,
        MILESTONE_TABLE="\n".join(f"- {m.name}: {m.status.value}" for m in milestones),
        TASK_TABLE=_task_table(tasks),
    )

    xlsx_dest = out_dir / f"{file_base}.xlsx"
    pstart = project_report_start_date(db, project)
    _, _, week_end = resolve_report_period(
        project.weekly_report_anchor_weekday,
        project.weekly_report_cutoff_offset_days,
        report.week_start,
        project_start_date=pstart,
        explicit_first_report_date=project.weekly_report_first_anchor_date,
    )
    frozen = report.frozen_metrics or {}
    planned = float(frozen.get("planned_pct") or 0)
    actual = float(frozen.get("actual_pct") or 0)
    spi_raw = frozen.get("spi")
    spi = float(spi_raw) if spi_raw is not None else None
    try:
        export_yyyymmdd_workbook(
            xlsx_dest,
            db=db,
            project=project,
            milestones=list(milestones),
            anchor=report.week_start,
            cut_off=week_end,
            planned_pct=planned,
            actual_pct=actual,
            spi=spi,
        )
    except FileNotFoundError:
        try:
            copy_template("weekly_report.xlsx", xlsx_dest)
            replace_in_xlsx(xlsx_dest, mapping)
            _append_tasks_sheet(xlsx_dest, tasks)
        except FileNotFoundError:
            xlsx_dest = _write_fallback_xlsx(project, report, milestones, out_dir, file_base)

    pptx_dest = out_dir / f"{file_base}.pptx"
    from app.models import ProjectPo, ProjectSph

    sph = db.get(ProjectSph, project.id)
    po = db.get(ProjectPo, project.id)
    try:
        build_weekly_report_pptx(
            pptx_dest,
            db=db,
            project=project,
            report=report,
            sph=sph,
            po=po,
        )
    except Exception:
        try:
            copy_template("weekly_report.pptx", pptx_dest)
            replace_in_pptx(pptx_dest, mapping)
        except FileNotFoundError:
            pptx_dest = None  # type: ignore

    return str(xlsx_dest), str(pptx_dest) if pptx_dest and pptx_dest.exists() else ""


def _append_tasks_sheet(path: Path, tasks) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(str(path))
    if "Tasks" in wb.sheetnames:
        ws = wb["Tasks"]
    else:
        ws = wb.create_sheet("Tasks")
    ws.append(["Task", "Status", "Due", "Spent ms", "Estimate ms"])
    for t in tasks:
        ws.append([t.name, t.status, str(t.due_date or ""), t.time_spent_ms, t.time_estimate_ms])
    wb.save(str(path))


def _write_fallback_xlsx(project, report, milestones, out_dir, file_base: str) -> Path:
    fname = f"{file_base}.xlsx"
    path = out_dir / fname
    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    ws.append(["Project", project.name])
    ws.append(["Week", str(report.week_start)])
    ws.append(["Planned %", report.frozen_metrics.get("planned_pct")])
    ws.append(["Actual %", report.frozen_metrics.get("actual_pct")])
    sm = report.summary or {}
    if sm.get("highlights"):
        ws.append(["Highlights", sm.get("highlights")])
    if sm.get("mitigation_plan"):
        ws.append(["Mitigasi mengejar plan", sm.get("mitigation_plan")])
    ms = wb.create_sheet("Milestone")
    ms.append(["Name", "Status", "Target", "Weight"])
    for m in milestones:
        ms.append([m.name, m.status.value, str(m.target_date or ""), m.weight_pct])
    wb.save(path)
    return path
