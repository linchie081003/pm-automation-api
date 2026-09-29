from datetime import date
from unittest.mock import patch

from openpyxl import Workbook

from app.services.yyyymmdd_report_excel import _fill_log_mingguan


def test_log_mingguan_clears_actual_after_cut_off():
    wb = Workbook()
    ws = wb.active
    ws.title = "Log_Mingguan"
    project = type(
        "P",
        (),
        {
            "id": 1,
            "weekly_report_anchor_weekday": 4,
            "weekly_report_cutoff_offset_days": 6,
        },
    )()
    anchors = [date(2026, 7, 2), date(2026, 7, 9), date(2026, 7, 16)]
    cut_off = date(2026, 7, 9)
    fake_actuals = {d: 0.42 for d in anchors}

    class FakeSession:
        def scalars(self, *_a, **_k):
            class R:
                def all(self):
                    return []

            return R()

    db = FakeSession()
    with (
        patch(
            "app.services.yyyymmdd_report_excel._weekly_report_actual_map",
            return_value=fake_actuals,
        ),
        patch("app.services.yyyymmdd_report_excel._snapshot_actual_map_exact", return_value={}),
    ):
        _fill_log_mingguan(
            ws,
            db=db,
            project=project,
            cut_off=cut_off,
            planned_pct=50,
            actual_pct=45,
            deviation_pp=-5,
            spi=0.9,
            spi_green_min=1.0,
            spi_yellow_min=0.9,
            anchors=anchors,
            project_start=date(2026, 6, 1),
            week_col_count=3,
        )
    assert ws.cell(8, 4).value == 0.42
    assert ws.cell(9, 4).value == 0.42
    assert ws.cell(10, 4).value is None
